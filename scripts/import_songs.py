#!/usr/bin/env python3
"""Validate an Excel song list and merge it into songs.json by Instagram URL."""
import argparse
from datetime import date, datetime
import json
from pathlib import Path
import sys
import tempfile
from urllib.parse import urlsplit

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
COLUMNS = {
    'title': 'title', 'artist': 'artist', 'title_en': 'en',
    'title_zh': 'zh', 'instagram_url': 'url',
    'post_date': 'date', 'source_work': 'work',
}
REQUIRED = {'title', 'artist', 'instagram_url'}


def text(value):
    return '' if value is None else str(value).strip()


def url_key(value):
    parsed = urlsplit(value)
    if (parsed.scheme != 'https' or parsed.hostname not in {'instagram.com', 'www.instagram.com'}
            or parsed.username or parsed.password or parsed.port
            or not parsed.path.startswith(('/reel/', '/p/', '/tv/'))
            or len(parsed.path.strip('/').split('/')) != 2):
        raise ValueError('instagram_url 必須是 HTTPS Instagram reel/p/tv 貼文連結')
    return 'https://www.instagram.com' + parsed.path.rstrip('/') + '/'


def post_date(value):
    if value is None or value == '':
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    value = text(value).replace('/', '-')
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        raise ValueError('post_date 必須是 Excel 日期儲存格或 YYYY-MM-DD 日期') from None


def read_excel(path, sheet):
    # Treat cell contents only as data; never evaluate formulas or external links.
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = iter(worksheet.iter_rows())
        header = next(rows, None)
        if header is None:
            raise ValueError('工作表是空的')
        headers = [text(cell.value) for cell in header]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError('欄位名稱重複')
        missing = REQUIRED - set(headers)
        if missing:
            raise ValueError('缺少必要欄位：' + ', '.join(sorted(missing)))
        imported = []
        seen = set()
        errors = []
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None for cell in cells):
                continue
            try:
                values = {headers[i]: cell.value for i, cell in enumerate(cells)
                          if i < len(headers) and headers[i] in COLUMNS}
                if any(cell.data_type == 'f' for i, cell in enumerate(cells)
                       if i < len(headers) and headers[i] in COLUMNS):
                    raise ValueError('匯入欄位不支援公式，請先貼上為值')
                for field in REQUIRED:
                    if not text(values.get(field)):
                        raise ValueError(f'{field} 不可空白')
                song = {target: text(values.get(source)) for source, target in COLUMNS.items()}
                song['url'] = url_key(song['url'])
                song['date'] = post_date(values.get('post_date'))
                if song['url'] in seen:
                    raise ValueError('Excel 中有重複的 Instagram 貼文')
                seen.add(song['url'])
                # Only supplied columns participate in updates; omitted columns retain existing data.
                supplied = {COLUMNS[source] for source in values}
                imported.append((song, supplied))
            except (ValueError, TypeError) as exc:
                errors.append(f'第 {number} 列：{exc}')
        if errors:
            raise ValueError('\n'.join(errors))
        if not imported:
            raise ValueError('沒有可匯入的歌曲')
        return imported
    finally:
        workbook.close()


def merge(existing, imported):
    if not isinstance(existing, list):
        raise ValueError('現有 songs.json 必須是陣列')
    merged = [dict(song) for song in existing]
    positions = {}
    for index, song in enumerate(merged):
        key = url_key(song['url'])
        if key in positions:
            raise ValueError('現有 songs.json 有重複貼文，請先處理')
        positions[key] = index
    added = updated = unchanged = 0
    for song, supplied in imported:
        key = song['url']
        if key in positions:
            index = positions[key]
            candidate = {**merged[index], **{field: song[field] for field in supplied}}
            if candidate == merged[index]:
                unchanged += 1
            else:
                merged[index] = candidate
                updated += 1
        else:
            positions[key] = len(merged)
            merged.append(song)
            added += 1
    return merged, added, updated, unchanged


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('excel', type=Path)
    parser.add_argument('--sheet', help='工作表名稱，預設第一張')
    parser.add_argument('--output', type=Path, default=ROOT / 'songs.json')
    parser.add_argument('--dry-run', action='store_true', help='只驗證與預覽，不寫入')
    args = parser.parse_args()
    try:
        imported = read_excel(args.excel, args.sheet)
        existing = json.loads(args.output.read_text(encoding='utf-8')) if args.output.exists() else []
        songs, added, updated, unchanged = merge(existing, imported)
        print(f'新增 {added} 首、更新 {updated} 首、未變更 {unchanged} 首；總計 {len(songs)} 首。')
        if not args.dry_run:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            temporary = None
            try:
                with tempfile.NamedTemporaryFile('w', encoding='utf-8', dir=args.output.parent,
                                                 suffix='.tmp', delete=False) as stream:
                    temporary = Path(stream.name)
                    json.dump(songs, stream, ensure_ascii=False, indent=2)
                    stream.write('\n')
                temporary.replace(args.output)
            finally:
                if temporary and temporary.exists():
                    temporary.unlink()
            print(f'已寫入 {args.output}')
    except (ValueError, KeyError, OSError) as exc:
        print(f'匯入失敗：{exc}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
